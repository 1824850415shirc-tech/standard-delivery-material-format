#!/usr/bin/env python3
"""Compare protected technical identifiers before and after editorial rewriting."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document


TOKEN_PATTERNS = [
    re.compile(r"https?://[^\s<>()]+", re.I),
    re.compile(r"(?<![\w.])/(?:[A-Za-z0-9_.-]+/)*[A-Za-z0-9_.-]+"),
    re.compile(r"\b[A-Za-z]:\\(?:[^\s\\]+\\)*[^\s\\]+"),
    re.compile(r"\b[A-Za-z0-9_.-]+\.(?:ya?ml|json|toml|py|sh|conf|ini|xml|onnx|txt|pdf|png|jpe?g|docx|wpt)\b", re.I),
    re.compile(r"\b[A-Z][A-Z0-9_]{2,}\b"),
    re.compile(r"(?<!\w)--?[A-Za-z][A-Za-z0-9_-]*"),
    re.compile(r"\b\d+\.\d+(?:\.\d+){0,2}\b"),
    re.compile(r"\b\d+(?:\.\d+)?%\b"),
]


def read_text(path: Path) -> str:
    if path.suffix.casefold() == ".docx":
        document = Document(path)
        chunks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                chunks.extend(cell.text for cell in row.cells)
        return "\n".join(chunks)
    return path.read_text(encoding="utf-8")


def tokens(text: str) -> set[str]:
    found: set[str] = set()
    for pattern in TOKEN_PATTERNS:
        for match in pattern.finditer(text):
            found.add(match.group(0).rstrip(".,;:，。；："))
    return found


def compare(before: Path, after: Path) -> dict:
    before_tokens = tokens(read_text(before))
    after_tokens = tokens(read_text(after))
    missing = sorted(before_tokens - after_tokens)
    added = sorted(after_tokens - before_tokens)
    return {
        "before": str(before),
        "after": str(after),
        "before_token_count": len(before_tokens),
        "after_token_count": len(after_tokens),
        "missing_from_after": missing,
        "added_in_after": added,
        "ok": not missing,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("before", type=Path)
    parser.add_argument("after", type=Path)
    parser.add_argument("--strict-added", action="store_true", help="Also fail when new identifiers appear.")
    args = parser.parse_args()
    report = compare(args.before, args.after)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    failed = bool(report["missing_from_after"]) or (args.strict_added and bool(report["added_in_after"]))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
