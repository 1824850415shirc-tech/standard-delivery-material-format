#!/usr/bin/env python3
"""Structurally verify a DOCX produced by the standard renderer."""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}
TARGET_FONTS = {"FangSong_GB2312", "仿宋_GB2312"}
HEADING_SIZES = {
    "文档标题": 15,
    "heading 1": 15,
    "heading 2": 15,
    "heading 3": 14,
    "heading 4": 14,
    "heading 5": 14,
}
TOC_STYLE_INDENTS = {
    "toc 1": 0,
    "toc 2": 420,
}
TOC_MAPPING = '\\o "1-2"'
PAGE_WIDTH_TWIPS = 11906
PAGE_HEIGHT_TWIPS = 16838
PAGE_MARGIN_TWIPS = 1440
TOC_RIGHT_TAB_TWIPS = PAGE_WIDTH_TWIPS - 2 * PAGE_MARGIN_TWIPS


def has_nonzero_indent(element) -> bool:
    indent = element.find(f".//{{{W_NS}}}ind")
    if indent is None:
        return False
    for name in ("left", "leftChars", "right", "rightChars", "firstLine", "firstLineChars", "hanging", "hangingChars"):
        value = indent.get(f"{{{W_NS}}}{name}")
        if value not in (None, "0"):
            return True
    return False


def toc_indent_matches(element, expected_left: int) -> bool:
    indent = element.find(f".//{{{W_NS}}}ind")
    if indent is None:
        return expected_left == 0
    left = indent.get(f"{{{W_NS}}}left")
    if int(left or 0) != expected_left:
        return False
    for name in ("leftChars", "right", "rightChars", "firstLine", "firstLineChars", "hanging", "hangingChars"):
        value = indent.get(f"{{{W_NS}}}{name}")
        if value not in (None, "0"):
            return False
    return True


def toc_tab_matches(element, expected_position: int) -> bool:
    tabs = element.findall(f".//{{{W_NS}}}tabs/{{{W_NS}}}tab")
    matching = [
        tab
        for tab in tabs
        if tab.get(f"{{{W_NS}}}val") == "right"
        and tab.get(f"{{{W_NS}}}leader") == "dot"
        and int(tab.get(f"{{{W_NS}}}pos") or 0) == expected_position
    ]
    return len(matching) == 1


def verify(path: Path, final: bool = False) -> dict:
    errors: list[str] = []
    warnings: list[str] = []
    document = Document(path)
    section_geometry: list[dict[str, int]] = []
    for index, section in enumerate(document.sections, start=1):
        geometry = {
            "page_width": int(round(section.page_width.twips)),
            "page_height": int(round(section.page_height.twips)),
            "margin_top": int(round(section.top_margin.twips)),
            "margin_bottom": int(round(section.bottom_margin.twips)),
            "margin_left": int(round(section.left_margin.twips)),
            "margin_right": int(round(section.right_margin.twips)),
        }
        section_geometry.append(geometry)
        expected_geometry = {
            "page_width": PAGE_WIDTH_TWIPS,
            "page_height": PAGE_HEIGHT_TWIPS,
            "margin_top": PAGE_MARGIN_TWIPS,
            "margin_bottom": PAGE_MARGIN_TWIPS,
            "margin_left": PAGE_MARGIN_TWIPS,
            "margin_right": PAGE_MARGIN_TWIPS,
        }
        if geometry != expected_geometry:
            errors.append(f"Section {index} must use A4 portrait with 25.4 mm margins: {geometry}")
    body_text = "\n".join(p.text for p in document.paragraphs)
    if "{{AI_CONTENT}}" in body_text:
        errors.append("Content marker remains in the document.")
    if re.search(r"(?m)^#{1,6}\s|```|~~~", body_text):
        errors.append("Raw Markdown syntax remains in document paragraphs.")
    if not body_text.strip() and not document.tables:
        errors.append("Document contains no generated content.")

    style_counts: dict[str, int] = {}
    page_break_indices: list[int] = []
    heading_indices: list[int] = []
    for paragraph in document.paragraphs:
        style_counts[paragraph.style.name] = style_counts.get(paragraph.style.name, 0) + 1
    for paragraph_index, paragraph in enumerate(document.paragraphs):
        if paragraph._p.xpath('.//w:br[@w:type="page"]'):
            page_break_indices.append(paragraph_index)
        if paragraph.style.name.casefold().startswith("heading ") or paragraph.style.name == "文档标题":
            heading_indices.append(paragraph_index)
            if any(run.bold is False for run in paragraph.runs):
                errors.append(f"Heading run explicitly disables bold formatting: {paragraph.text[:60]}")

    styles_by_name = {style.name.casefold(): style for style in document.styles}
    for style_name, expected_size in HEADING_SIZES.items():
        style = styles_by_name.get(style_name)
        if style is None:
            errors.append(f"Missing heading style: {style_name}")
            continue
        if style.font.bold is not True:
            errors.append(f"Heading style is not explicitly bold: {style_name}")
        if style.font.size is None or abs(style.font.size.pt - expected_size) > 0.01:
            errors.append(f"Heading style has wrong size: {style_name}; expected {expected_size} pt")
        if has_nonzero_indent(style.element):
            errors.append(f"Heading style has a nonzero paragraph indent: {style_name}")

    table_style = styles_by_name.get("表格正文".casefold())
    if table_style is None or table_style.font.size is None or abs(table_style.font.size.pt - 14) > 0.01:
        errors.append("表格正文 must use 四号 (14 pt).")

    toc_paragraphs = [p for p in document.paragraphs if p.style.name.casefold().startswith("toc ")]
    toc_level_counts: dict[str, int] = {}
    for paragraph in toc_paragraphs:
        name = paragraph.style.name.casefold()
        toc_level_counts[name] = toc_level_counts.get(name, 0) + 1
    if final:
        if not toc_paragraphs:
            errors.append("Final verification requires a refreshed TOC.")
        for paragraph in toc_paragraphs:
            name = paragraph.style.name.casefold()
            if name not in TOC_STYLE_INDENTS:
                errors.append(f"TOC contains an unsupported level: {paragraph.style.name}")
                continue
            expected = TOC_STYLE_INDENTS[name]
            if not toc_indent_matches(paragraph.style.element, expected) or not toc_indent_matches(paragraph._p, expected):
                errors.append(
                    f"TOC entry has the wrong indent for {paragraph.style.name}; "
                    f"expected {expected} twips: {paragraph.text[:60]}"
                )
            if not toc_tab_matches(paragraph.style.element, TOC_RIGHT_TAB_TWIPS) or not toc_tab_matches(
                paragraph._p, TOC_RIGHT_TAB_TWIPS
            ):
                errors.append(
                    f"TOC entry has the wrong right tab for {paragraph.style.name}; "
                    f"expected {TOC_RIGHT_TAB_TWIPS} twips: {paragraph.text[:60]}"
                )

    if len(page_break_indices) < 2:
        errors.append("Front matter must contain page breaks after the cover and after the TOC.")
    if "{{COVER_COMPANY}}" in body_text:
        errors.append("Cover company placeholder remains in the document.")
    if heading_indices and len(page_break_indices) >= 2 and heading_indices[0] <= page_break_indices[1]:
        errors.append("Generated body headings appear before the TOC page break.")
    table_paragraphs = 0
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    table_paragraphs += 1
                    if paragraph.style.name != "表格正文":
                        errors.append(f"Table paragraph does not use 表格正文: {paragraph.text[:40]}")

    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        required = {"word/document.xml", "word/styles.xml", "word/numbering.xml"}
        missing = sorted(required - names)
        if missing:
            errors.append("Missing required package parts: " + ", ".join(missing))
        styles_root = etree.fromstring(package.read("word/styles.xml"))
        document_root = etree.fromstring(package.read("word/document.xml"))
        toc_instructions = document_root.xpath('//w:instrText[contains(., "TOC")]/text()', namespaces=NS)
        if not toc_instructions:
            errors.append("Document has no real TOC field.")
        elif not any(TOC_MAPPING in value for value in toc_instructions):
            errors.append("TOC field must include Heading 1 through Heading 2 only.")
        for style_name in ["Normal", "文档标题", "heading 1", "heading 2", "heading 3", "heading 4", "heading 5", "表格正文"]:
            style = styles_root.xpath(f"//w:style[w:name/@w:val='{style_name}']", namespaces=NS)
            if not style:
                errors.append(f"Missing template style: {style_name}")
                continue
            fonts = style[0].xpath(".//w:rFonts", namespaces=NS)
            if not fonts:
                errors.append(f"Style has no explicit font mapping: {style_name}")
                continue
            values = {value for value in fonts[0].attrib.values()}
            if not values.intersection(TARGET_FONTS):
                errors.append(f"Style does not reference target FangSong font: {style_name}")

    if not any(name.casefold().startswith("heading ") for name in style_counts):
        warnings.append("No heading paragraphs were generated.")
    return {
        "input": str(path),
        "section_geometry": section_geometry,
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "table_paragraphs": table_paragraphs,
        "page_break_indices": page_break_indices,
        "first_heading_index": heading_indices[0] if heading_indices else None,
        "toc_paragraphs": len(toc_paragraphs),
        "toc_level_counts": toc_level_counts,
        "final_mode": final,
        "style_counts": style_counts,
        "errors": errors,
        "warnings": warnings,
        "ok": not errors,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--json", action="store_true")
    parser.add_argument("--final", action="store_true", help="Require a refreshed two-level TOC with the standard indents.")
    args = parser.parse_args()
    report = verify(args.input, final=args.final)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
