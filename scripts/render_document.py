#!/usr/bin/env python3
"""Render compliant Markdown into the retained WPS-derived DOCX template."""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import unquote, urlparse

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.parts.numbering import NumberingPart
from docx.shared import Cm, Pt, Twips


SKILL_DIR = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "master-template.docx"
MARKER = "{{AI_CONTENT}}"
FONT_ASCII = "FangSong_GB2312"
FONT_EAST_ASIA = "仿宋_GB2312"
TOC_FIELD_INSTRUCTION = ' TOC \\o "1-2" \\h \\z '
DOCUMENT_TITLE_STYLE = "文档标题"
PAGE_WIDTH_TWIPS = 11906
PAGE_HEIGHT_TWIPS = 16838
PAGE_MARGIN_TWIPS = 1440
TOC_STYLE_INDENTS = {
    "toc 1": 0,
    "toc 2": 420,
}
STYLE_MAP = {
    1: DOCUMENT_TITLE_STYLE,
    2: "heading 1",
    3: "heading 2",
    4: "heading 3",
    5: "heading 4",
}
HEADING_SIZES = {
    DOCUMENT_TITLE_STYLE: 15,
    "heading 1": 15,
    "heading 2": 15,
    "heading 3": 14,
    "heading 4": 14,
    "heading 5": 14,
}
INLINE_RE = re.compile(
    r"(`[^`]+`|\*\*.+?\*\*|__.+?__|(?<!\*)\*[^*]+\*(?!\*)|(?<!_)_[^_]+_(?!_)|\[[^\]]+\]\([^)]+\))"
)
CHECKBOX_RE = re.compile(r"^\[([ xX])\]\s+(.+)$")
IMAGE_RE = re.compile(r"^!\[([^\]]*)\]\((.+)\)$")
LOCAL_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg"}
EMU_PER_TWIP = 635
MAX_IMAGE_HEIGHT_RATIO = 0.72


@dataclass
class Block:
    kind: str
    data: object
    extra: object = None


def parse_frontmatter(lines: list[str]) -> tuple[dict[str, str], list[str]]:
    if not lines or lines[0].strip() != "---":
        return {}, lines
    end = next((i for i in range(1, len(lines)) if lines[i].strip() == "---"), None)
    if end is None:
        return {}, lines
    metadata: dict[str, str] = {}
    for line in lines[1:end]:
        if ":" in line:
            key, value = line.split(":", 1)
            metadata[key.strip()] = value.strip().strip('"\'')
    return metadata, lines[end + 1 :]


def is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def starts_block(lines: list[str], index: int) -> bool:
    stripped = lines[index].strip()
    if not stripped:
        return True
    if stripped.startswith("<!--"):
        return True
    if IMAGE_RE.match(stripped):
        return True
    if re.match(r"^#{1,6}\s+", stripped):
        return True
    if re.match(r"^(```+|~~~+)", stripped):
        return True
    if re.match(r"^\s*(?:[-+*]|\d+[.)])\s+", lines[index]):
        return True
    if stripped.startswith(">") or re.fullmatch(r"(?:-{3,}|\*{3,}|_{3,})", stripped):
        return True
    if "|" in stripped and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
        return True
    return False


def parse_image_target(payload: str) -> tuple[str, str | None]:
    """Parse a local Markdown image destination and optional quoted caption."""
    value = payload.strip()
    caption = None
    caption_match = re.search(r'\s+"([^"]*)"\s*$', value)
    if caption_match:
        caption = caption_match.group(1).strip() or None
        value = value[: caption_match.start()].strip()
    if value.startswith("<") and value.endswith(">"):
        value = value[1:-1].strip()
    if not value:
        raise ValueError("Markdown image path must not be empty.")
    return unquote(value), caption


def parse_markdown(text: str) -> tuple[dict[str, str], list[Block]]:
    metadata, lines = parse_frontmatter(text.splitlines())
    blocks: list[Block] = []
    list_group = 0
    i = 0
    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("<!--"):
            while i < len(lines) and "-->" not in lines[i]:
                i += 1
            i += 1
            continue
        image_match = IMAGE_RE.match(stripped)
        if image_match:
            alt_text = image_match.group(1).strip()
            if not alt_text:
                raise ValueError("Markdown images require meaningful alt text.")
            destination, caption = parse_image_target(image_match.group(2))
            blocks.append(Block("image", destination, (alt_text, caption)))
            i += 1
            continue
        fence = re.match(r"^(```+|~~~+)\s*([^\s]*)", stripped)
        if fence:
            marker = fence.group(1)
            language = fence.group(2)
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith(marker[0] * 3):
                code.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            blocks.append(Block("code", "\n".join(code), language))
            continue
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", stripped)
        if heading:
            blocks.append(Block("heading", heading.group(2), min(len(heading.group(1)), 5)))
            i += 1
            continue
        if "|" in stripped and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            rows = [split_table_row(raw)]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(split_table_row(lines[i]))
                i += 1
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            blocks.append(Block("table", rows))
            continue
        list_match = re.match(r"^(\s*)([-+*]|\d+[.)])\s+(.+)$", raw)
        if list_match:
            if not blocks or blocks[-1].kind != "list_item":
                list_group += 1
            marker = list_match.group(2)
            content = list_match.group(3)
            checkbox = CHECKBOX_RE.match(content) if not marker[0].isdigit() else None
            if checkbox:
                kind = "checklist_checked" if checkbox.group(1).casefold() == "x" else "checklist_unchecked"
                content = checkbox.group(2)
            else:
                kind = "ordered" if marker[0].isdigit() else "bullet"
            level = min(len(list_match.group(1).expandtabs(4)) // 2, 8)
            blocks.append(Block("list_item", content, (kind, level, list_group)))
            i += 1
            continue
        if stripped.startswith(">"):
            quote_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            quote_text = " ".join(quote_lines)
            explanation = re.match(r"^说明\s*[：:]\s*(.+)$", quote_text)
            if explanation:
                blocks.append(Block("paragraph", explanation.group(1)))
            else:
                blocks.append(Block("quote", quote_text))
            continue
        if re.fullmatch(r"(?:-{3,}|\*{3,}|_{3,})", stripped):
            blocks.append(Block("rule", ""))
            i += 1
            continue
        paragraph = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not starts_block(lines, i):
            paragraph.append(lines[i].strip())
            i += 1
        blocks.append(Block("paragraph", " ".join(paragraph)))
    return metadata, blocks


def find_style(document: Document, name: str):
    for style in document.styles:
        if style.name.casefold() == name.casefold():
            return style
    raise KeyError(f"Template style not found: {name}")


def set_rfonts(rpr, ascii_name: str = FONT_ASCII, east_asia: str = FONT_EAST_ASIA) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key, value in {
        "w:ascii": ascii_name,
        "w:hAnsi": ascii_name,
        "w:eastAsia": east_asia,
        "w:cs": ascii_name,
    }.items():
        rfonts.set(qn(key), value)


def set_indent(element, left_twips: int = 0, first_line_twips: int = 0) -> None:
    ppr = element.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:left"), str(left_twips))
    ind.set(qn("w:firstLine"), str(first_line_twips))
    for attribute in ("leftChars", "right", "rightChars", "firstLineChars"):
        ind.set(qn(f"w:{attribute}"), "0")
    for attribute in ("hanging", "hangingChars"):
        ind.attrib.pop(qn(f"w:{attribute}"), None)


def set_zero_indent(element) -> None:
    set_indent(element, 0, 0)


def ensure_page_geometry(document: Document) -> None:
    """Apply the standard A4 portrait page with uniform 25.4 mm margins."""
    for section in document.sections:
        section.page_width = Twips(PAGE_WIDTH_TWIPS)
        section.page_height = Twips(PAGE_HEIGHT_TWIPS)
        section.top_margin = Twips(PAGE_MARGIN_TWIPS)
        section.bottom_margin = Twips(PAGE_MARGIN_TWIPS)
        section.left_margin = Twips(PAGE_MARGIN_TWIPS)
        section.right_margin = Twips(PAGE_MARGIN_TWIPS)


def available_width_twips(document: Document) -> int:
    section = document.sections[0]
    return int(round(section.page_width.twips - section.left_margin.twips - section.right_margin.twips))


def set_toc_right_tab(element, position_twips: int) -> None:
    """Replace TOC tab stops with one right-aligned dotted page-number tab."""
    ppr = element.get_or_add_pPr()
    tabs = ppr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        ppr.append(tabs)
    for tab in list(tabs):
        tabs.remove(tab)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), str(position_twips))
    tabs.append(tab)


def configure_toc_field(document: Document) -> None:
    instructions = document.element.xpath('.//w:instrText[contains(., "TOC")]')
    if not instructions:
        raise RuntimeError("Template has no real TOC field.")
    instructions[0].text = TOC_FIELD_INSTRUCTION
    for extra in instructions[1:]:
        if "TOC" in (extra.text or ""):
            extra.text = ""


def ensure_document_title_style(document: Document) -> None:
    normal = find_style(document, "Normal")
    try:
        style = find_style(document, DOCUMENT_TITLE_STYLE)
    except KeyError:
        style = document.styles.add_style(DOCUMENT_TITLE_STYLE, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = normal
    set_rfonts(style.element.get_or_add_rPr())
    style.font.size = Pt(15)
    style.font.bold = True
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.keep_together = True
    set_zero_indent(style.element)
    ppr = style.element.get_or_add_pPr()
    outline = ppr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        ppr.append(outline)
    outline.set(qn("w:val"), "9")


def ensure_toc_styles(document: Document) -> None:
    normal = find_style(document, "Normal")
    right_tab = available_width_twips(document)
    for name, left_twips in TOC_STYLE_INDENTS.items():
        try:
            style = find_style(document, name)
        except KeyError:
            style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = normal
        set_rfonts(style.element.get_or_add_rPr())
        style.font.size = Pt(14)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        set_indent(style.element, left_twips, 0)
        set_toc_right_tab(style.element, right_tab)


def normalize_toc_paragraphs(document: Document) -> None:
    right_tab = available_width_twips(document)
    for paragraph in document.paragraphs:
        name = paragraph.style.name.casefold()
        if name not in TOC_STYLE_INDENTS:
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_indent(paragraph._p, TOC_STYLE_INDENTS[name], 0)
        set_toc_right_tab(paragraph._p, right_tab)


def normalize_template_fonts(document: Document) -> None:
    ensure_page_geometry(document)
    configure_toc_field(document)
    ensure_document_title_style(document)
    for style_name in ["Normal", *HEADING_SIZES, "表格正文"]:
        style = find_style(document, style_name)
        set_rfonts(style.element.get_or_add_rPr())
    for style_name, size in HEADING_SIZES.items():
        style = find_style(document, style_name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        set_zero_indent(style.element)
    table_body = find_style(document, "表格正文")
    table_body.font.size = Pt(14)
    table_body.paragraph_format.line_spacing = 1
    set_zero_indent(table_body.element)
    ensure_toc_styles(document)
    normalize_toc_paragraphs(document)
    root = document.styles.element
    rpr = root.find("./" + qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
    if rpr is not None:
        set_rfonts(rpr)


def set_optional_bold(run, bold: bool | None) -> None:
    if bold is not None:
        run.bold = bold


def add_inline(paragraph, text: str, bold: bool | None = None) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            set_optional_bold(paragraph.add_run(text[cursor : match.start()]), bold)
        token = match.group(0)
        run = None
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "EDEDED")
            run._r.get_or_add_rPr().append(shading)
        elif token.startswith("**") or token.startswith("__"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") or token.startswith("_"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("["):
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            run = paragraph.add_run(f"{link.group(1)}（{link.group(2)}）" if link else token)
        if run is not None and bold is True:
            run.bold = True
        cursor = match.end()
    if cursor < len(text):
        set_optional_bold(paragraph.add_run(text[cursor:]), bold)


class NumberingManager:
    def __init__(self, document: Document):
        try:
            self.part = document.part.numbering_part
        except (KeyError, NotImplementedError):
            element = parse_xml(f"<w:numbering {nsdecls('w')}/>")
            self.part = NumberingPart(
                PackURI("/word/numbering.xml"),
                CT.WML_NUMBERING,
                element,
                document.part.package,
            )
            document.part.relate_to(self.part, RT.NUMBERING)
        self.root = self.part.element
        self.next_abstract = 1 + max(
            [int(e.get(qn("w:abstractNumId"))) for e in self.root.findall(qn("w:abstractNum"))] or [-1]
        )
        self.next_num = 1 + max(
            [int(e.get(qn("w:numId"))) for e in self.root.findall(qn("w:num"))] or [0]
        )
        self.abstract = {
            "bullet": self._add_abstract("bullet"),
            "ordered": self._add_abstract("ordered"),
            "checklist_unchecked": self._add_abstract("checklist_unchecked"),
            "checklist_checked": self._add_abstract("checklist_checked"),
        }

    def _add_abstract(self, kind: str) -> int:
        abstract_id = self.next_abstract
        self.next_abstract += 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "multilevel")
        abstract.append(multi)
        for level in range(9):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(level))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            num_fmt = OxmlElement("w:numFmt")
            num_fmt.set(qn("w:val"), "decimal" if kind == "ordered" else "bullet")
            lvl_text = OxmlElement("w:lvlText")
            if kind == "bullet":
                lvl_text.set(qn("w:val"), "•")
            elif kind == "checklist_unchecked":
                lvl_text.set(qn("w:val"), "□")
            elif kind == "checklist_checked":
                lvl_text.set(qn("w:val"), "☒")
            else:
                lvl_text.set(qn("w:val"), ".".join(f"%{n}" for n in range(1, level + 2)) + ".")
            suffix = OxmlElement("w:suff")
            suffix.set(qn("w:val"), "space")
            lvl_jc = OxmlElement("w:lvlJc")
            lvl_jc.set(qn("w:val"), "left")
            ppr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            left = 480 + level * 360
            ind.set(qn("w:left"), str(left))
            ind.set(qn("w:hanging"), "240")
            ppr.append(ind)
            rpr = OxmlElement("w:rPr")
            set_rfonts(rpr)
            for child in (start, num_fmt, lvl_text, suffix, lvl_jc, ppr, rpr):
                lvl.append(child)
            abstract.append(lvl)
        self.root.append(abstract)
        return abstract_id

    def new_list(self, kind: str) -> int:
        num_id = self.next_num
        self.next_num += 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_id = OxmlElement("w:abstractNumId")
        abstract_id.set(qn("w:val"), str(self.abstract[kind]))
        num.append(abstract_id)
        self.root.append(num)
        return num_id

    @staticmethod
    def apply(paragraph, num_id: int, level: int) -> None:
        ppr = paragraph._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(level))
        num = OxmlElement("w:numId")
        num.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num])
        ppr.append(num_pr)


def add_paragraph_before(document: Document, marker, style_name: str, text: str = ""):
    paragraph = document.add_paragraph()
    paragraph.style = find_style(document, style_name)
    if text:
        add_inline(paragraph, text)
    marker._p.addprevious(paragraph._p)
    return paragraph


def resolve_local_image_path(markdown_path: Path, destination: str) -> Path:
    parsed = urlparse(destination)
    if parsed.scheme in {"http", "https", "data"} or destination.startswith("//"):
        raise ValueError(f"Remote Markdown images are not supported: {destination}")
    if parsed.scheme == "file":
        image_path = Path(unquote(parsed.path))
    elif parsed.scheme:
        raise ValueError(f"Unsupported Markdown image source: {destination}")
    else:
        image_path = Path(destination)
        if not image_path.is_absolute():
            image_path = markdown_path.parent / image_path
    image_path = image_path.resolve()
    if image_path.suffix.casefold() not in LOCAL_IMAGE_SUFFIXES:
        allowed = ", ".join(sorted(LOCAL_IMAGE_SUFFIXES))
        raise ValueError(f"Unsupported image format for {image_path}; expected one of: {allowed}")
    if not image_path.is_file():
        raise FileNotFoundError(f"Markdown image not found: {image_path}")
    return image_path


def add_image_before(document: Document, marker, image_path: Path, alt_text: str, caption: str | None) -> None:
    paragraph = add_paragraph_before(document, marker, "Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_after = Pt(0)
    shape = paragraph.add_run().add_picture(str(image_path))

    section = document.sections[0]
    maximum_width = available_width_twips(document) * EMU_PER_TWIP
    body_height_twips = int(round(section.page_height.twips - section.top_margin.twips - section.bottom_margin.twips))
    maximum_height = int(body_height_twips * MAX_IMAGE_HEIGHT_RATIO * EMU_PER_TWIP)
    scale = min(1.0, maximum_width / shape.width, maximum_height / shape.height)
    if scale < 1.0:
        shape.width = int(shape.width * scale)
        shape.height = int(shape.height * scale)

    properties = shape._inline.docPr
    properties.set("descr", alt_text)
    properties.set("title", caption or alt_text)

    if caption:
        caption_paragraph = add_paragraph_before(document, marker, "Normal", caption)
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.paragraph_format.first_line_indent = Cm(0)
        caption_paragraph.paragraph_format.space_before = Pt(0)
        caption_paragraph.paragraph_format.space_after = Pt(6)


def set_paragraph_shading(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def set_quote_border(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "808080")
    borders.append(left)
    ppr.append(borders)


def display_width(text: str) -> int:
    return sum(2 if unicodedata.east_asian_width(ch) in {"W", "F"} else 1 for ch in text)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "000000")


def set_table_geometry(table, widths: list[int], total_width: int) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total_width))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    grid_cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
    for index, width in enumerate(widths):
        if index < len(grid_cols):
            grid_cols[index].set(qn("w:w"), str(width))
        for cell in table.columns[index].cells:
            cell.width = Twips(width)
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def add_table_before(document: Document, marker, rows: list[list[str]]):
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = find_style(document, "Normal Table")
    set_table_borders(table)
    available_dxa = available_width_twips(document)
    weights = []
    for col in range(column_count):
        width = max(display_width(row[col]) if col < len(row) else 0 for row in rows)
        weights.append(max(8, min(width, 48)))
    minimum = 900
    remaining = max(0, available_dxa - minimum * column_count)
    total_weight = sum(weights) or column_count
    widths = [minimum + int(remaining * weight / total_weight) for weight in weights]
    widths[-1] += available_dxa - sum(widths)
    set_table_geometry(table, widths, available_dxa)

    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.style = find_style(document, "表格正文")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            value = values[col_index] if col_index < len(values) else ""
            add_inline(paragraph, value, bold=row_index == 0)
    marker._p.addprevious(table._tbl)
    return table


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def replace_metadata(document: Document, metadata: dict[str, str], document_title: str) -> None:
    resolved = {
        "title": metadata.get("title") or document_title,
        "project_name": metadata.get("project_name") or "【待补充：项目名称】",
        "cover_title": metadata.get("cover_title") or metadata.get("title") or document_title,
        "cover_company": metadata.get("cover_company") or "XX科技",
        "document_date": metadata.get("document_date") or "【待补充：日期】",
        "document_version": metadata.get("document_version") or "V【待补充：版本号】",
    }
    resolved.update(metadata)
    values = {f"{{{{{key.upper()}}}}}": value for key, value in resolved.items()}
    values.update({
        "XX项目": resolved["project_name"],
        "技术规格说明书": resolved["cover_title"],
        "XX科技": resolved["cover_company"],
        "XX年X月X日": resolved["document_date"],
    })
    document.core_properties.title = resolved["title"]
    document.core_properties.author = resolved["cover_company"]
    document.core_properties.last_modified_by = resolved["cover_company"]

    def replace_in_paragraph(paragraph) -> None:
        for run in paragraph.runs:
            for placeholder, value in values.items():
                if placeholder in run.text:
                    run.text = run.text.replace(placeholder, value)

    for paragraph in document.paragraphs:
        replace_in_paragraph(paragraph)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph)
    for section in document.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                replace_in_paragraph(paragraph)


def set_update_fields(document: Document) -> None:
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def render(input_path: Path, template_path: Path, output_path: Path) -> dict:
    metadata, blocks = parse_markdown(input_path.read_text(encoding="utf-8"))
    document = Document(template_path)
    normalize_template_fonts(document)
    document_title = next(
        (str(block.data) for block in blocks if block.kind == "heading" and int(block.extra) == 1),
        metadata.get("title") or "【待补充：文档名称】",
    )
    replace_metadata(document, metadata, document_title)
    marker = next((p for p in document.paragraphs if p.text.strip() == MARKER), None)
    if marker is None:
        raise RuntimeError(f"Template content marker not found: {MARKER}")

    numbering = NumberingManager(document)
    list_ids: dict[tuple[str, int], int] = {}
    counts = {
        "paragraph": 0,
        "heading": 0,
        "list_item": 0,
        "quote": 0,
        "code": 0,
        "table": 0,
        "image": 0,
    }

    for block in blocks:
        if block.kind == "heading":
            paragraph = add_paragraph_before(document, marker, STYLE_MAP[int(block.extra)], str(block.data))
            counts["heading"] += 1
        elif block.kind == "paragraph":
            add_paragraph_before(document, marker, "Normal", str(block.data))
            counts["paragraph"] += 1
        elif block.kind == "list_item":
            kind, level, group = block.extra
            key = (kind, group)
            if key not in list_ids:
                list_ids[key] = numbering.new_list(kind)
            paragraph = add_paragraph_before(document, marker, "Normal", str(block.data))
            paragraph.paragraph_format.first_line_indent = None
            NumberingManager.apply(paragraph, list_ids[key], level)
            counts["list_item"] += 1
        elif block.kind == "quote":
            paragraph = add_paragraph_before(document, marker, "Normal", str(block.data))
            paragraph.paragraph_format.left_indent = Cm(0.74)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            set_quote_border(paragraph)
            counts["quote"] += 1
        elif block.kind == "code":
            paragraph = add_paragraph_before(document, marker, "Normal")
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1
            run = paragraph.add_run(str(block.data))
            set_paragraph_shading(paragraph, "F2F2F2")
            counts["code"] += 1
        elif block.kind == "table":
            add_table_before(document, marker, block.data)
            counts["table"] += 1
        elif block.kind == "image":
            alt_text, caption = block.extra
            image_path = resolve_local_image_path(input_path, str(block.data))
            add_image_before(document, marker, image_path, alt_text, caption)
            counts["image"] += 1
        elif block.kind == "rule":
            paragraph = add_paragraph_before(document, marker, "Normal")
            paragraph.paragraph_format.first_line_indent = Cm(0)
            ppr = paragraph._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:color"), "808080")
            borders.append(bottom)
            ppr.append(borders)

    remove_paragraph(marker)
    set_update_fields(document)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return {
        "input": str(input_path),
        "template": str(template_path),
        "output": str(output_path),
        "blocks": counts,
        "font_ascii": FONT_ASCII,
        "font_east_asia": FONT_EAST_ASIA,
        "toc_mapping": "Markdown H2 -> Heading 1 -> TOC 1; Markdown H3 -> Heading 2 -> TOC 2",
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    report = render(args.input, args.template, args.output)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
