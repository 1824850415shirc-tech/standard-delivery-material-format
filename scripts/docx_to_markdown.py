#!/usr/bin/env python3
"""Extract ordinary DOCX paragraphs and tables into standardizable Markdown."""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


@dataclass
class Item:
    kind: str
    value: object
    level: int | None = None


def iter_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def paragraph_heading_level(paragraph: Paragraph) -> int | None:
    name = paragraph.style.name.casefold().strip()
    match = re.fullmatch(r"heading\s*([1-5])", name)
    if match:
        return int(match.group(1))
    if name in {"title", "标题"}:
        return 1
    if name == "文档标题":
        return 1
    return None


def is_numbered(paragraph: Paragraph) -> bool:
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.numPr is not None


def clean_cell(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().replace("|", "\\|")


def table_markdown(table: Table) -> str:
    rows = [[clean_cell(cell.text) for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    if len(rows) == 1:
        rows.append([""] * width)
    lines = ["| " + " | ".join(rows[0]) + " |"]
    lines.append("| " + " | ".join(["---"] * width) + " |")
    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n".join(lines)


def extract(input_path: Path, output_path: Path, title: str | None = None) -> dict:
    document = Document(input_path)
    items: list[Item] = []
    skipped = {"front_matter": 0, "cover": 0, "toc": 0, "empty": 0}
    has_standard_front_matter = any(p.style.name.casefold() == "封面" for p in document.paragraphs) and any(
        p.style.name.casefold().startswith("toc ") for p in document.paragraphs
    )
    body_started = not has_standard_front_matter

    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            text = re.sub(r"\s+", " ", block.text).strip()
            style_name = block.style.name.casefold().strip()
            level = paragraph_heading_level(block)
            if not body_started:
                if level:
                    body_started = True
                else:
                    skipped["front_matter"] += 1
                    continue
            if style_name == "封面":
                skipped["cover"] += 1
                continue
            if style_name.startswith("toc ") or text.replace(" ", "") == "目录":
                skipped["toc"] += 1
                continue
            if not text or text == "{{AI_CONTENT}}":
                skipped["empty"] += 1
                continue
            if level:
                items.append(Item("heading", text, level))
            elif is_numbered(block):
                items.append(Item("list", text))
            else:
                items.append(Item("paragraph", text))
        else:
            if not body_started:
                skipped["front_matter"] += 1
                continue
            markdown = table_markdown(block)
            if markdown:
                items.append(Item("table", markdown))

    heading_items = [item for item in items if item.kind == "heading"]
    h1_items = [item for item in heading_items if item.level == 1]
    first_heading = next((item for item in items if item.kind == "heading"), None)
    has_document_title = len(h1_items) == 1 and first_heading is h1_items[0]
    inferred_title = title or document.core_properties.title or input_path.stem
    if not has_document_title:
        shallowest = min((item.level or 1) for item in heading_items) if heading_items else 2
        shift = 2 - shallowest
        for item in heading_items:
            item.level = max(2, min((item.level or 1) + shift, 5))
        items.insert(0, Item("heading", inferred_title, 1))

    lines: list[str] = []
    for item in items:
        if item.kind == "heading":
            lines.append("#" * int(item.level or 1) + " " + str(item.value))
        elif item.kind == "list":
            lines.append("- " + str(item.value))
        else:
            lines.append(str(item.value))
        lines.append("")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    image_count = len(document.inline_shapes)
    warnings: list[str] = []
    if image_count:
        warnings.append("Inline images were not embedded in Markdown; preserve or reinsert them during DOCX editing.")
    if len(document.sections) > 1:
        warnings.append("The source has multiple sections; rebuilding from Markdown will not preserve section-specific layout.")
    return {
        "input": str(input_path),
        "output": str(output_path),
        "items": {
            "headings": sum(item.kind == "heading" for item in items),
            "paragraphs": sum(item.kind == "paragraph" for item in items),
            "lists": sum(item.kind == "list" for item in items),
            "tables": sum(item.kind == "table" for item in items),
        },
        "skipped": skipped,
        "inline_images": image_count,
        "sections": len(document.sections),
        "warnings": warnings,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--title")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    report = extract(args.input, args.output, args.title)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
