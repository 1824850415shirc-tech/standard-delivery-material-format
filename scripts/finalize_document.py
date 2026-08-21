#!/usr/bin/env python3
"""Finalize heading, TOC, and table styles after WPS/Word refreshes fields."""

from __future__ import annotations

import argparse
import os
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from render_document import (
    FONT_ASCII,
    FONT_EAST_ASIA,
    TOC_STYLE_INDENTS,
    find_style,
    normalize_template_fonts,
    available_width_twips,
    set_indent,
    set_rfonts,
    set_toc_right_tab,
)


TOC_STYLES = set(TOC_STYLE_INDENTS)


def finalize(input_path: Path, output_path: Path) -> None:
    document = Document(input_path)
    normalize_template_fonts(document)
    toc_right_tab = available_width_twips(document)

    for paragraph in document.paragraphs:
        if paragraph.style.name.casefold() in TOC_STYLES:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            left_twips = TOC_STYLE_INDENTS[paragraph.style.name.casefold()]
            set_indent(paragraph._p, left_twips, 0)
            set_indent(paragraph.style.element, left_twips, 0)
            set_toc_right_tab(paragraph._p, toc_right_tab)
            set_toc_right_tab(paragraph.style.element, toc_right_tab)
            set_rfonts(paragraph.style.element.get_or_add_rPr(), FONT_ASCII, FONT_EAST_ASIA)
            paragraph.style.font.size = Pt(14)

    table_style = find_style(document, "表格正文")
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.style = table_style

    output_path.parent.mkdir(parents=True, exist_ok=True)
    if input_path.resolve() == output_path.resolve():
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, dir=output_path.parent) as handle:
            temporary = Path(handle.name)
        try:
            document.save(temporary)
            os.replace(temporary, output_path)
        finally:
            temporary.unlink(missing_ok=True)
    else:
        document.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    finalize(args.input, args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
