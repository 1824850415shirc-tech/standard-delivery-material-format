#!/usr/bin/env python3
"""Regression checks for cover metadata, checklist spacing, and callout handling."""

from __future__ import annotations

import struct
import sys
import tempfile
import unittest
import zipfile
import zlib
from pathlib import Path

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from render_document import available_width_twips, render, set_indent, set_toc_right_tab  # noqa: E402
from verify_document import verify  # noqa: E402


NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def write_test_png(path: Path, width: int = 640, height: int = 360) -> None:
    def chunk(kind: bytes, payload: bytes) -> bytes:
        checksum = zlib.crc32(kind + payload) & 0xFFFFFFFF
        return struct.pack(">I", len(payload)) + kind + payload + struct.pack(">I", checksum)

    row = b"\x00" + b"\xDB\xEA\xF5" * width
    payload = b"".join(
        [b"\x89PNG\r\n\x1a\n", chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))]
        + [chunk(b"IDAT", zlib.compress(row * height)), chunk(b"IEND", b"")]
    )
    path.write_bytes(payload)


class RenderRegressionTest(unittest.TestCase):
    def test_cover_checklists_and_callouts(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "regression.docx"
            render(
                ROOT / "tests" / "fixtures" / "render-regression.md",
                ROOT / "assets" / "master-template.docx",
                output,
            )

            report = verify(output)
            self.assertTrue(report["ok"], report["errors"])

            document = Document(output)
            body_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("XX科技", body_text)
            self.assertNotIn("这里只用于指导写作", body_text)
            self.assertEqual(document.core_properties.author, "XX科技")
            self.assertEqual(document.core_properties.last_modified_by, "XX科技")

            checklist_paragraphs = [
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text.startswith(("已确认维护窗口", "已核对配置文件"))
            ]
            self.assertEqual(len(checklist_paragraphs), 2)
            for paragraph in checklist_paragraphs:
                self.assertNotIn("[ ]", paragraph.text)
                self.assertFalse(paragraph._p.xpath("./w:pPr/w:pBdr"))
                self.assertTrue(paragraph._p.xpath("./w:pPr/w:numPr"))

            explanation = next(
                paragraph for paragraph in document.paragraphs if paragraph.text.startswith("这段内容会转成普通正文")
            )
            self.assertFalse(explanation.text.startswith("说明："))
            self.assertFalse(explanation._p.xpath("./w:pPr/w:pBdr"))

            warning = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("注意："))
            self.assertTrue(warning._p.xpath("./w:pPr/w:pBdr"))

            with zipfile.ZipFile(output) as package:
                numbering = etree.fromstring(package.read("word/numbering.xml"))
            checkbox_levels = numbering.xpath("//w:lvl[w:lvlText/@w:val='□']", namespaces=NS)
            self.assertTrue(checkbox_levels)
            level = checkbox_levels[0]
            self.assertEqual(level.xpath("string(w:suff/@w:val)", namespaces=NS), "space")
            self.assertEqual(level.xpath("string(w:pPr/w:ind/@w:left)", namespaces=NS), "480")
            self.assertEqual(level.xpath("string(w:pPr/w:ind/@w:hanging)", namespaces=NS), "240")

    def test_local_image_caption_alt_text_and_relative_path(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            image_path = root / "登录 页面.png"
            markdown_path = root / "manual.md"
            output = root / "manual.docx"
            write_test_png(image_path)
            markdown_path.write_text(
                "---\n"
                "cover_title: 操作手册\n"
                "cover_company: XX科技\n"
                "project_name: 【待补充：系统名称】\n"
                "document_version: V0.1\n"
                "document_date: 2026-09-02\n"
                "---\n\n"
                "# 操作手册\n\n"
                "## 1. 开始使用\n\n"
                "### 1.1 登录\n\n"
                "![登录页面](<登录 页面.png> \"图 1-1 登录页面\")\n",
                encoding="utf-8",
            )

            report = render(markdown_path, ROOT / "assets" / "master-template.docx", output)
            self.assertEqual(report["blocks"]["image"], 1)
            verification = verify(output)
            self.assertTrue(verification["ok"], verification["errors"])
            self.assertEqual(verification["images"], 1)

            document = Document(output)
            self.assertEqual(len(document.inline_shapes), 1)
            self.assertIn("图 1-1 登录页面", [paragraph.text for paragraph in document.paragraphs])
            with zipfile.ZipFile(output) as package:
                document_xml = etree.fromstring(package.read("word/document.xml"))
            descriptions = document_xml.xpath(
                "//wp:docPr/@descr",
                namespaces={"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"},
            )
            self.assertEqual(descriptions, ["登录页面"])

    def test_remote_images_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            markdown_path = root / "manual.md"
            markdown_path.write_text(
                "# 操作手册\n\n## 1. 开始使用\n\n![登录页面](https://example.com/login.png)\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "Remote Markdown images are not supported"):
                render(markdown_path, ROOT / "assets" / "master-template.docx", root / "manual.docx")

    def test_final_verification_accepts_wps_effective_toc_formatting(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "wps-toc.docx"
            render(
                ROOT / "tests" / "fixtures" / "render-regression.md",
                ROOT / "assets" / "master-template.docx",
                output,
            )

            document = Document(output)
            for style_name, left_twips in (("toc 1", 0), ("toc 2", 420)):
                paragraph = document.add_paragraph(f"{style_name} 回归条目\t3", style=style_name)
                set_indent(paragraph._p, left_twips, 0)
                set_toc_right_tab(paragraph._p, available_width_twips(document))
            toc_paragraphs = [
                paragraph for paragraph in document.paragraphs if paragraph.style.name.casefold().startswith("toc ")
            ]
            self.assertTrue(toc_paragraphs)
            for paragraph in toc_paragraphs:
                direct_indent = paragraph._p.find("./w:pPr/w:ind", namespaces=NS)
                if direct_indent is not None:
                    direct_indent.getparent().remove(direct_indent)
                if paragraph.style.name.casefold() == "toc 2":
                    style_indent = paragraph.style.element.find("./w:pPr/w:ind", namespaces=NS)
                    self.assertIsNotNone(style_indent)
                    style_indent.set(f"{{{NS['w']}}}leftChars", "200")
                style_tabs = paragraph.style.element.find("./w:pPr/w:tabs", namespaces=NS)
                if style_tabs is not None:
                    style_tabs.getparent().remove(style_tabs)
            document.save(output)

            report = verify(output, final=True)
            self.assertTrue(report["ok"], report["errors"])


if __name__ == "__main__":
    unittest.main()
